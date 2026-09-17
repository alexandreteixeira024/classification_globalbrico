"""Testes unitários para a sincronização incremental de emails e atualização do Excel."""

import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.email_data import LABELS, load_labelled_emails
from src.sync_excel import (
    canonical_label,
    copy_new_files_from_source,
    get_or_create_workbook,
    normalize_email,
    sync_emails,
)


class TestSyncExcel(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.source_dir = self.root / "source"
        self.input_dir = self.root / "incoming"
        self.excel_path = self.root / "test_classificacao.xlsx"

        self.source_dir.mkdir()
        self.input_dir.mkdir()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_canonical_label(self):
        self.assertEqual(canonical_label("Pedido de Informação"), "Pedido de Informação")
        self.assertEqual(canonical_label("pedido de informação"), "Pedido de Informação")
        self.assertEqual(canonical_label("pedido de informacao"), "Pedido de Informação")
        self.assertEqual(canonical_label("spam"), "SPAM")
        self.assertEqual(canonical_label("Pedido de Encomenda"), "Pedido de Encomenda")
        self.assertEqual(canonical_label(""), "")
        self.assertEqual(canonical_label("Outra Categoria"), "Outra Categoria")

    def test_normalize_email_various_formats(self):
        # Formato IMAP
        data1 = {
            "uid": "1001",
            "from": "cliente@example.com",
            "subject": "Dúvida sobre produto",
            "date": "Wed, 16 Oct 2024 10:00:00 +0100",
            "text": "Gostaria de saber o preço.",
            "attachments": [{"filename": "catalogo.pdf"}],
        }
        norm1 = normalize_email(data1)
        self.assertEqual(norm1["uid"], "1001")
        self.assertEqual(norm1["from"], "cliente@example.com")
        self.assertEqual(norm1["subject"], "Dúvida sobre produto")
        self.assertIn("catalogo.pdf", norm1["attachments_desc"])

        # Formato API com id e outra IA
        data2 = {
            "id": "ENC-9999",
            "sender": "fornecedor@empresa.com",
            "subject": "Encomenda 456",
            "received_at": "2026-09-17T12:00:00Z",
            "body": "Segue a nossa encomenda.",
            "outra_ia": "Pedido de Encomenda",
        }
        norm2 = normalize_email(data2)
        self.assertEqual(norm2["uid"], "ENC-9999")
        self.assertEqual(norm2["from"], "fornecedor@empresa.com")
        self.assertEqual(norm2["outra_ia_label"], "Pedido de Encomenda")

    def test_copy_new_files_from_source(self):
        # Cria ficheiros na pasta externa
        (self.source_dir / "email1.json").write_text('{"uid": "1"}', encoding="utf-8")
        (self.source_dir / "email2.json").write_text('{"uid": "2"}', encoding="utf-8")
        (self.source_dir / "summary.json").write_text('{}', encoding="utf-8")

        copied = copy_new_files_from_source(self.source_dir, self.input_dir)
        self.assertEqual(copied, 2)
        self.assertTrue((self.input_dir / "email1.json").exists())
        self.assertTrue((self.input_dir / "email2.json").exists())
        self.assertFalse((self.input_dir / "summary.json").exists())

        # Segunda cópia não deve duplicar
        copied_again = copy_new_files_from_source(self.source_dir, self.input_dir)
        self.assertEqual(copied_again, 0)

    def test_sync_emails_incremental(self):
        # 1. Cria 2 emails iniciais
        email_a = {
            "uid": "101",
            "subject": "Orçamento para ferramentas",
            "from": "compras@obra.pt",
            "date": "2026-09-15",
            "text": "Podem enviar cotação?",
            "outra_ia": "Pedido de Informação",
        }
        email_b = {
            "uid": "102",
            "subject": "***SPAM*** Ganhe um prémio",
            "from": "promo@spam.com",
            "date": "2026-09-15",
            "text": "Clique aqui para ganhar.",
            "outra_ia": "SPAM",
        }
        (self.input_dir / "101.json").write_text(json.dumps(email_a), encoding="utf-8")
        (self.input_dir / "102.json").write_text(json.dumps(email_b), encoding="utf-8")

        # Primeira sincronização
        added = sync_emails(
            input_dir=self.input_dir,
            excel_path=self.excel_path,
            predict_transformer=False,
        )
        self.assertEqual(added, 2)
        self.assertTrue(self.excel_path.exists())

        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Revisão"]
        self.assertEqual(ws.max_row, 3)  # Cabeçalho + 2 linhas

        # Simula preenchimento manual da 'Label correta' na linha 2
        header = [cell.value for cell in ws[1]]
        uid_col = header.index("UID") + 1
        correct_col = header.index("Label correta") + 1
        ws.cell(row=2, column=correct_col, value="Pedido de Informação")
        wb.save(self.excel_path)

        # 2. Re-execução sem novos ficheiros
        added_empty = sync_emails(
            input_dir=self.input_dir,
            excel_path=self.excel_path,
            predict_transformer=False,
        )
        self.assertEqual(added_empty, 0)

        # Verifica se o preenchimento manual continua lá intacto
        wb_check = openpyxl.load_workbook(self.excel_path)
        ws_check = wb_check["Revisão"]
        self.assertEqual(ws_check.cell(row=2, column=correct_col).value, "Pedido de Informação")

        # 3. Adiciona um terceiro email
        email_c = {
            "uid": "103",
            "subject": "Confirmação de Encomenda 789",
            "from": "cliente@loja.pt",
            "date": "2026-09-17",
            "text": "Confirmamos o envio da encomenda.",
            "outra_ia": "Pedido de Encomenda",
        }
        (self.input_dir / "103.json").write_text(json.dumps(email_c), encoding="utf-8")

        added_c = sync_emails(
            input_dir=self.input_dir,
            excel_path=self.excel_path,
            predict_transformer=False,
        )
        self.assertEqual(added_c, 1)

        wb_final = openpyxl.load_workbook(self.excel_path)
        ws_final = wb_final["Revisão"]
        self.assertEqual(ws_final.max_row, 4)  # Cabeçalho + 3 linhas

        # Verifica se a linha 2 mantém a label humana e se a nova linha 4 foi inserida
        self.assertEqual(ws_final.cell(row=2, column=correct_col).value, "Pedido de Informação")
        self.assertEqual(ws_final.cell(row=4, column=uid_col).value, "103")

    def test_compatibility_with_load_labelled_emails(self):
        # Garante que load_labelled_emails lê perfeitamente este Excel
        email_data = {
            "uid": "2001",
            "subject": "Pedido urgente",
            "from": "geral@empresa.pt",
            "date": "2026-09-17",
            "text": "Necessitamos de cotação urgente.",
        }
        (self.input_dir / "2001.json").write_text(json.dumps(email_data), encoding="utf-8")

        sync_emails(
            input_dir=self.input_dir,
            excel_path=self.excel_path,
            predict_transformer=False,
        )

        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Revisão"]
        header = [cell.value for cell in ws[1]]
        correct_col = header.index("Label correta") + 1
        ws.cell(row=2, column=correct_col, value="Pedido de Informação")
        wb.save(self.excel_path)

        # Chama a função existente no projeto
        labelled = load_labelled_emails(self.excel_path, self.input_dir)
        self.assertEqual(len(labelled), 1)
        self.assertEqual(labelled[0]["uid"], "2001")
        self.assertEqual(labelled[0]["label"], "Pedido de Informação")


if __name__ == "__main__":
    unittest.main()
