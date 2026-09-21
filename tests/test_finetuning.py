import json
import tempfile
import unittest
from collections import Counter
from pathlib import Path

import openpyxl

from src.finetuning import load_examples, stratified_folds
from src.sync_excel import sync_emails


LABELS = ("Pedido de Informação", "Pedido de Encomenda", "SPAM")


def examples_per_class(count=5):
    return [
        (f"{label_index}-{item_index}", f"texto {label_index}-{item_index}", label)
        for label_index, label in enumerate(LABELS)
        for item_index in range(count)
    ]


class StratifiedFoldTests(unittest.TestCase):
    def test_each_email_is_evaluated_once_and_folds_are_stratified(self):
        examples = examples_per_class()
        folds = stratified_folds(examples, folds=5, seed=42)

        observed = [item[0] for fold in folds for item in fold]
        self.assertCountEqual(observed, [item[0] for item in examples])
        self.assertEqual(len(observed), len(set(observed)))
        for fold in folds:
            self.assertEqual(Counter(item[2] for item in fold), Counter(LABELS))

    def test_load_examples_ignores_replies_already_present_in_excel(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            emails = root / "emails"
            emails.mkdir()
            (emails / "initial.json").write_text(json.dumps({
                "uid": "initial", "subject": "Novo pedido", "text": "Texto",
                "in_reply_to": None,
            }), encoding="utf-8")
            (emails / "reply.json").write_text(json.dumps({
                "uid": "reply", "subject": "Re: Novo pedido", "text": "Resposta",
                "in_reply_to": "<initial@example.com>",
            }), encoding="utf-8")
            excel = root / "labels.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Revisão"
            sheet.append(["UID", "Label correta"])
            sheet.append(["initial", LABELS[0]])
            sheet.append(["reply", LABELS[1]])
            workbook.save(excel)

            loaded = load_examples(excel, [emails])

        self.assertEqual([item[0] for item in loaded], ["initial"])

    def test_sync_excel_does_not_insert_replies(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            emails = root / "emails"
            emails.mkdir()
            for uid, in_reply_to in (("initial", None), ("reply", "<initial@example.com>")):
                (emails / f"{uid}.json").write_text(json.dumps({
                    "uid": uid, "subject": uid, "text": "Texto",
                    "in_reply_to": in_reply_to,
                }), encoding="utf-8")
            excel = root / "classification.xlsx"

            added = sync_emails(
                input_dir=emails,
                excel_path=excel,
                predict_transformer=False,
            )
            workbook = openpyxl.load_workbook(excel, read_only=True)
            values = list(workbook["Revisão"].iter_rows(values_only=True))
            workbook.close()

        self.assertEqual(added, 1)
        self.assertEqual([row[0] for row in values[1:]], ["initial"])


if __name__ == "__main__":
    unittest.main()
